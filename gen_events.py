import openpyxl, json, re, sys, os

SRC = sys.argv[1] if len(sys.argv)>1 else '/root/.claude/uploads/7845ef9d-bcab-5ced-bd75-80ccafdb5a95/aca5e2d9-THJW_Activity_Visibility_Form_Responses.xlsx'

# Display fixes for typos in the form's own option lists.
FIX = {
  'Democracy & Authoritarism': 'Democracy & Authoritarianism',
  'Policmakers/Diplomats': 'Policymakers/Diplomats',
  'Visual arts- Culture and Politics': 'Visual Arts, Culture & Politics',
}
def fix(s): return FIX.get(s.strip(), s.strip())

PERSONISH = re.compile(r'^[A-Z][a-zÀ-ſ\'\-]+ [A-Z][a-zÀ-ſ\'\-]+$')
ORG_HINT = re.compile(r'museum|institute|foundation|university|centre|center|council|association|ministry|embassy|federation|committee|organisation|organization|network|society|academy|chambers|court|hub|law|trust|initiative|mission|school', re.I)

def host_of(org, owner):
    org = (org or '').strip()
    if org: return org
    owner = (owner or '').strip()
    if owner and not PERSONISH.match(owner) and ORG_HINT.search(owner):
        return owner
    return ''

def split_multi(v):
    return [fix(x) for x in re.split(r'\s*,\s*(?![^(]*\))', (v or '').strip()) if x.strip()]

URL = re.compile(r'https?://\S+')

wb = openpyxl.load_workbook(SRC)
ws = wb[wb.sheetnames[0]]
hdr = [ (c.value or '').strip() for c in ws[1] ]
def col(name, default=None):
    for i,h in enumerate(hdr):
        if h.lower().startswith(name.lower()): return i
    return default

C = {k: col(k) for k in ['Timestamp','Activity Owner Name','Activity Owner E-mail','Activity Title',
     'Activity Description','Target Audience','Activity Format','Activity Registration','Activity Tags',
     'Activity Public-Facing Contact','Name of Organisation']}
# optional columns that will appear once the form is extended
C['Activity Date'] = col('Activity Date')
C['Start Time']    = col('Start Time')
C['End Time']      = col('End Time')
C['Venue']         = col('Venue') or col('Location')

def get(row, key):
    i = C.get(key)
    if i is None or i >= len(row): return ''
    v = row[i]
    return '' if v is None else (v if not hasattr(v,'strftime') else v)

events=[]
for n,row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
    if not any(row): continue
    title = str(get(row,'Activity Title') or '').strip()
    if not title: continue
    reg = str(get(row,'Activity Registration') or '').strip()
    m = URL.search(reg)
    d = get(row,'Activity Date')
    date = d.strftime('%Y-%m-%d') if hasattr(d,'strftime') else str(d or '').strip()
    events.append({
      'id': 'e%02d' % n,
      'title': title,
      'host': host_of(str(get(row,'Name of Organisation') or ''), str(get(row,'Activity Owner Name') or '')),
      'date': date,
      'start': str(get(row,'Start Time') or '').strip(),
      'end': str(get(row,'End Time') or '').strip(),
      'venue': str(get(row,'Venue') or '').strip(),
      'format': fix(str(get(row,'Activity Format') or '').strip()),
      'tags': split_multi(str(get(row,'Activity Tags') or '')),
      'audience': split_multi(str(get(row,'Target Audience') or '')),
      'description': re.sub(r'\s+\n','\n', str(get(row,'Activity Description') or '').strip()),
      'register': m.group(0).rstrip('.,);') if m else '',
      'contact': str(get(row,'Activity Public-Facing Contact') or '').strip(),
    })

out={'updated': __import__('datetime').date.today().isoformat(), 'events': events}
json.dump(out, open('events.json','w'), indent=1, ensure_ascii=False)
print(len(events),'events ->  events.json')
for e in events:
    print(' -', (e['host'] or '(host missing)')[:34].ljust(34), '|', e['date'] or 'no date', '|',
          e['format'], '|', ('link' if e['register'] else 'no link'), '|', e['title'][:44])
