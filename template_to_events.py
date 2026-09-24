#!/usr/bin/env python3
"""Turn a filled-in THJW activity template into events.json for the programme widget.

    python3 template_to_events.py "THJW-2026-activity-template.xlsx" [more.xlsx ...]

Several files can be passed at once: colleagues can each return their own copy and the
activities are merged into one programme, sorted by date. Rows still carrying the grey
example are skipped, as are rows with no title.

The programme is rebuilt from scratch on every run, so pass every file, not just the new one.
If the result would hold fewer activities than the current events.json, nothing is written
and the missing titles are listed. Add --allow-fewer when removing activities on purpose.
"""
import sys, json, re, datetime
from openpyxl import load_workbook

EXAMPLE_TITLE = "Justice has no Passport: A Global Conversation on Universal Jurisdiction"
EXAMPLE_ORG   = "Baltasar Garzón International Foundation"

HEAD = {  # template header -> field
 "organisation name":"host", "activity title":"title", "date":"date", "end date":"end_date",
 "start time":"start", "end time":"end", "venue or address":"venue",
 "format":"format", "topic 1":"t1", "topic 2":"t2", "topic 3":"t3",
 "audience 1":"a1", "audience 2":"a2", "audience 3":"a3",
 "activity description":"description", "registration link":"register",
 "public contact":"contact",
}
URL = re.compile(r"https?://\S+")

def txt(v):
    if v is None: return ""
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.date): return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.time): return v.strftime("%H:%M")
    return str(v).strip()

def end_date(start, end):
    """Keep an end date only when it falls after the start: a multi-day activity."""
    return end if start and end > start else ""

def read(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Activities"] if "Activities" in wb.sheetnames else wb[wb.sheetnames[0]]
    # find the header row (the one containing "Activity title")
    hrow = None
    for r in range(1, 12):
        vals = [txt(c.value).lower().rstrip(" *") for c in ws[r]]
        if "activity title" in vals:
            hrow = r; cols = vals; break
    if hrow is None:
        raise SystemExit("%s: could not find the header row" % path)
    idx = {}
    for i, name in enumerate(cols):
        if name in HEAD: idx[HEAD[name]] = i
    out = []
    for row in ws.iter_rows(min_row=hrow+1, values_only=True):
        g = lambda k: txt(row[idx[k]]) if k in idx and idx[k] < len(row) else ""
        title = g("title")
        if not title: continue
        if title == EXAMPLE_TITLE and g("host").startswith(EXAMPLE_ORG[:20]):
            continue                                   # untouched example row
        reg = g("register"); m = URL.search(reg)
        out.append({
          "title": title,
          "host": g("host"),
          "date": g("date"),
          "end_date": end_date(g("date"), g("end_date")),
          "start": g("start"),
          "end": g("end"),
          "venue": g("venue"),
          "format": g("format"),
          "tags": [x for x in (g("t1"), g("t2"), g("t3")) if x],
          "audience": [x for x in (g("a1"), g("a2"), g("a3")) if x],
          "description": g("description"),
          "register": m.group(0).rstrip(".,);") if m else "",
          "contact": g("contact"),
        })
    return out

def dropped(events):
    """Titles in the current events.json that the new programme would lose."""
    try:
        old = json.load(open("events.json", encoding="utf-8"))["events"]
    except (OSError, ValueError, KeyError):
        return None
    if len(events) >= len(old): return None
    new = {e["title"] for e in events}
    return [e["title"] for e in old if e.get("title") not in new]

def main(paths, allow_fewer=False):
    events = []
    for p in paths: events += read(p)
    gone = dropped(events)
    if gone is not None and not allow_fewer:
        print("Not written: this would publish %d activities, fewer than the current programme."
              % len(events))
        for t in gone: print("  would disappear: %s" % t[:70])
        print("Did you pass every file? Run it on ~/Downloads/thjw/*.xlsx.\n"
              "If you are removing activities on purpose, add --allow-fewer.")
        raise SystemExit(1)
    events.sort(key=lambda e: (e["date"] == "", e["date"], e["start"], e["title"]))
    for i, e in enumerate(events, 1): e["id"] = "e%02d" % i
    order = ["id","title","host","date","end_date","start","end","venue","format","tags","audience",
             "description","register","contact"]
    events = [{k: e[k] for k in order} for e in events]
    json.dump({"updated": datetime.date.today().isoformat(), "events": events},
              open("events.json", "w"), indent=1, ensure_ascii=False)
    print("%d activities -> events.json" % len(events))
    missing = [e["title"][:48] for e in events if not e["date"]]
    if missing:
        print("  no date yet (%d): %s" % (len(missing), "; ".join(missing)))
    nolink = [e["title"][:48] for e in events if not e["register"]]
    if nolink:
        print("  no registration link (%d): %s" % (len(nolink), "; ".join(nolink)))

if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--allow-fewer"]
    if not args: raise SystemExit(__doc__)
    main(args, allow_fewer=len(args) < len(sys.argv) - 1)
